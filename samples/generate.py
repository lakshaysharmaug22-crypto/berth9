#!/usr/bin/env python3
"""
Generates Berth 9 demo data: partner files in their native formats, the buyer's reference data
(open purchase orders, catalog, ship-to locations, FX) and the SQL seed the server loads.

Everything here is fictional. Some records are deliberately wrong (price variance, over-invoiced
quantity, bad GSTIN check digit, duplicate line, unknown SKU, truncated trailer count) so the
exceptions queue has something real to show.

    python3 samples/generate.py
"""
import json
import os
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INBOUND = os.path.join(ROOT, "samples", "inbound")
RESOURCES = os.path.join(ROOT, "berth9-server", "src", "main", "resources")
BERTH9 = os.path.join(RESOURCES, "berth9")

GSTIN_CHARS = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"


def gstin(first14):
    """Appends the mod-36 check character to the first 14 characters of a GSTIN."""
    total = 0
    for i, ch in enumerate(first14):
        product = GSTIN_CHARS.index(ch) * (1 if i % 2 == 0 else 2)
        total += product // 36 + product % 36
    return first14 + GSTIN_CHARS[(36 - total % 36) % 36]


def money(value):
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def inr(value):
    """Indian digit grouping: 1,23,45,678.50"""
    value = money(value)
    whole, frac = f"{value:.2f}".split(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        whole = ",".join(groups + [tail])
    return f"{whole}.{frac}"


SAHYADRI_GSTIN = gstin("27AAKCS4417R1Z")
NARMADA_GSTIN = gstin("24AAFFN8823K1Z")
NIMBUS_GSTIN = gstin("29AAGCN5521M1Z")
BUYER_GSTIN = gstin("27AAHCH6624Q1Z")

# ------------------------------------------------------------------ reference data (the buyer's systems)

PARTNERS = [
    {"id": "KESTREL", "name": "Kestrel Fasteners Inc.", "country": "US", "city": "Columbus, OH", "role": "SUPPLIER",
     "channel": "SFTP", "format": "X12", "document": "EDI 810 invoice", "schema": "invoice-line", "currency": "USD",
     "ediId": "KESTRELFAST", "locale": "US", "expectedBy": "10:00"},
    {"id": "LAKESHORE", "name": "Lakeshore Packaging Co.", "country": "US", "city": "Milwaukee, WI", "role": "SUPPLIER",
     "channel": "API", "format": "XML", "document": "cXML invoice", "schema": "invoice-line", "currency": "USD",
     "locale": "US", "expectedBy": "12:00"},
    {"id": "SAHYADRI", "name": "Sahyadri Polymers Pvt. Ltd.", "country": "IN", "city": "Pune, MH", "role": "SUPPLIER",
     "channel": "SFTP", "format": "XLSX", "document": "Excel invoice register", "schema": "invoice-line", "currency": "INR",
     "gstin": SAHYADRI_GSTIN, "locale": "IN", "expectedBy": "11:30"},
    {"id": "NARMADA", "name": "Narmada Metal Works", "country": "IN", "city": "Vadodara, GJ", "role": "SUPPLIER",
     "channel": "FOLDER", "format": "CSV", "document": "CSV invoice lines", "schema": "invoice-line", "currency": "INR",
     "gstin": NARMADA_GSTIN, "locale": "IN", "expectedBy": "13:00"},
    {"id": "MIDWEST", "name": "Midwest Bearing Supply", "country": "US", "city": "Des Moines, IA", "role": "SUPPLIER",
     "channel": "SFTP", "format": "FIXED_WIDTH", "document": "AS/400 fixed-width export", "schema": "invoice-line",
     "currency": "USD", "locale": "US", "expectedBy": "09:00"},
    {"id": "NIMBUS", "name": "Nimbus Components Pvt. Ltd.", "country": "IN", "city": "Bengaluru, KA", "role": "SUPPLIER",
     "channel": "API", "format": "JSON", "document": "JSON invoice (REST push)", "schema": "invoice-line", "currency": "USD",
     "gstin": NIMBUS_GSTIN, "locale": "IN", "expectedBy": "15:00"},
    {"id": "BRIGHTMART", "name": "Brightmart Retail", "country": "US", "city": "Dallas, TX", "role": "CUSTOMER",
     "channel": "SFTP", "format": "X12", "document": "EDI 850 purchase order", "schema": "order-line", "currency": "USD",
     "ediId": "BRIGHTMART", "locale": "US", "expectedBy": "08:00"},
]

PURCHASE_ORDERS = [
    ("PO-2026-00412", "KESTREL", "USD", "2026-09-02", [
        ("KF-HEX-M8-40", "Hex bolt M8x40 zinc", 5000, "0.42", "EA"),
        ("KF-HEX-M10-50", "Hex bolt M10x50 zinc", 3000, "0.61", "EA"),
        ("KF-NUT-M8", "Hex nut M8 zinc", 8000, "0.09", "EA"),
        ("KF-NUT-M10", "Hex nut M10 zinc", 6000, "0.12", "EA"),
        ("KF-WSH-M8", "Flat washer M8", 10000, "0.03", "EA"),
    ]),
    ("PO-2026-00437", "KESTREL", "USD", "2026-09-08", [
        ("KF-ANC-12", "Wedge anchor 1/2 x 4-1/4", 1200, "0.88", "EA"),
        ("KF-SCR-DW-2", "Drywall screw #8 x 2, box/1000", 400, "21.50", "BX"),
        ("KF-LAG-38", "Lag screw 3/8 x 3", 2500, "0.34", "EA"),
    ]),
    ("PO-2026-00455", "LAKESHORE", "USD", "2026-09-05", [
        ("LP-BOX-RSC-18", "Corrugated RSC box 18x12x12", 6000, "1.08", "EA"),
        ("LP-BOX-RSC-24", "Corrugated RSC box 24x18x18", 3000, "1.94", "EA"),
        ("LP-TAPE-48", "Carton sealing tape 48mm x 100m", 1440, "2.35", "RL"),
        ("LP-WRAP-18", "Stretch wrap 18in x 1500ft", 480, "17.90", "RL"),
        ("LP-VOID-FILL", "Kraft void-fill paper roll", 120, "38.00", "RL"),
        ("LP-LBL-4X6", "Thermal labels 4x6, roll/500", 600, "9.40", "RL"),
    ]),
    ("HLIN-PO-2627-0118", "SAHYADRI", "INR", "2026-09-01", [
        ("SP-HDPE-GRN-25", "HDPE granules natural, 25 kg bag", 400, "2850.00", "BG"),
        ("SP-PP-GRN-25", "PP granules, 25 kg bag", 300, "2640.00", "BG"),
        ("SP-LLDPE-FLM-500", "LLDPE film roll 500 mm", 200, "1980.00", "RL"),
        ("SP-MB-BLK-25", "Black masterbatch, 25 kg bag", 60, "4150.00", "BG"),
    ]),
    ("HLIN-PO-2627-0126", "SAHYADRI", "INR", "2026-09-04", [
        ("SP-HDPE-GRN-25", "HDPE granules natural, 25 kg bag", 250, "2850.00", "BG"),
        ("SP-PVC-CMP-25", "PVC compound, 25 kg bag", 150, "3420.00", "BG"),
    ]),
    ("HLIN-PO-2627-0144", "NARMADA", "INR", "2026-08-28", [
        ("NMW-FLG-SS304-2", "SS304 slip-on flange 2in", 400, "685.00", "EA"),
        ("NMW-FLG-SS304-3", "SS304 slip-on flange 3in", 250, "1140.00", "EA"),
        ("NMW-ELB-SS304-2", "SS304 90 deg elbow 2in", 600, "312.00", "EA"),
        ("NMW-TEE-SS304-2", "SS304 equal tee 2in", 300, "468.00", "EA"),
        ("NMW-PIPE-SS304-2-6M", "SS304 pipe 2in sch40, 6 m", 120, "9850.00", "EA"),
    ]),
    ("PO-2026-00468", "MIDWEST", "USD", "2026-09-03", [
        ("MBS-6204-2RS", "Deep groove ball bearing 6204-2RS", 2000, "2.18", "EA"),
        ("MBS-6205-2RS", "Deep groove ball bearing 6205-2RS", 1500, "2.64", "EA"),
        ("MBS-UCP205", "Pillow block bearing UCP205", 400, "11.75", "EA"),
        ("MBS-SEAL-TC-35", "Oil seal TC 35x52x7", 3000, "0.58", "EA"),
        ("MBS-GREASE-EP2", "EP2 grease cartridge 14 oz", 480, "4.95", "EA"),
    ]),
    ("HLIN-PO-2627-0160", "NIMBUS", "USD", "2026-09-06", [
        ("NC-PCB-CTRL-V3", "Controller PCB assembly v3", 500, "18.40", "EA"),
        ("NC-HARN-12P", "12-pin wiring harness, 1 m", 1500, "3.25", "EA"),
        ("NC-SENS-TEMP-K", "K-type thermocouple sensor", 800, "6.10", "EA"),
        ("NC-ENCL-IP65", "IP65 ABS enclosure 200x150", 400, "7.80", "EA"),
    ]),
]

CATALOG = [
    ("HL-GLV-NIT-M", "Nitrile gloves, medium, box/100", "BX", "12.50"),
    ("HL-GLV-NIT-L", "Nitrile gloves, large, box/100", "BX", "12.50"),
    ("HL-TAPE-48-CLR", "Carton tape 48 mm clear, case/36", "CS", "58.00"),
    ("HL-BOX-RSC-18", "RSC box 18x12x12, bundle/25", "BD", "34.50"),
    ("HL-LBL-4X6-500", "Thermal label roll 4x6", "RL", "14.25"),
    ("HL-STRAP-PP-12", "PP strapping 1/2 in coil", "RL", "44.00"),
    ("HL-BRG-6204", "Ball bearing 6204-2RS", "EA", "3.95"),
    ("HL-FLG-SS304-2", "SS304 flange 2 in", "EA", "14.80"),
]

SHIP_TO = [
    ("BRIGHTMART", "DC04", "Brightmart DC 04", "Dallas", "TX"),
    ("BRIGHTMART", "DC07", "Brightmart DC 07", "Houston", "TX"),
    ("BRIGHTMART", "DC11", "Brightmart DC 11", "Oklahoma City", "OK"),
]

FX = [("USD", "1"), ("INR", "0.0114"), ("EUR", "1.08"), ("GBP", "1.27")]


def write(path, content, mode="w"):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, mode, newline="" if mode == "w" else None) as f:
        f.write(content)


# ------------------------------------------------------------------ X12 helpers


def isa(sender, receiver, date6, time4, control):
    return ("ISA*00*          *00*          *ZZ*{:<15}*ZZ*{:<15}*{}*{}*U*00401*{:09d}*0*P*>"
            .format(sender, receiver, date6, time4, control))


def x12(segments):
    return "~\n".join(segments) + "~\n"


def transaction(st_control, type_, body):
    segs = [f"ST*{type_}*{st_control}"] + body
    segs.append(f"SE*{len(segs) + 1}*{st_control}")
    return segs


# ------------------------------------------------------------------ partner files


def kestrel():
    inv1 = [
        ("1", 1500, "EA", "0.42", "KF-HEX-M8-40", "HEX BOLT M8X40 ZINC"),
        ("2", 1000, "EA", "0.61", "KF-HEX-M10-50", "HEX BOLT M10X50 ZINC"),
        ("3", 2500, "EA", "0.09", "KF-NUT-M8", "HEX NUT M8 ZINC"),
        ("4", 2000, "EA", "0.125", "KF-NUT-M10", "HEX NUT M10 ZINC"),          # +4.2% vs PO: warning
        ("5", 5000, "EA", "0.03", "KF-WSH-M8", "FLAT WASHER M8"),
    ]
    inv2 = [
        ("1", 400, "EA", "0.88", "KF-ANC-12", "WEDGE ANCHOR 1/2 X 4-1/4"),
        ("2", 150, "BX", "21.50", "KF-SCR-DW-2", "DRYWALL SCREW #8 X 2 BOX/1000"),
        ("3", 600, "EA", "0.34", "KF-LAG-516", "LAG SCREW 5/16 X 3"),         # not on the PO: error
    ]

    def body(inv_no, inv_date, po_no, po_date, lines):
        segs = [f"BIG*{inv_date}*{inv_no}*{po_date}*{po_no}",
                "REF*IA*HRB-0042",
                "N1*SE*KESTREL FASTENERS INC*92*KESTREL",
                "N1*BY*HARBORLINE SUPPLY CO*92*HARBORLINE",
                "CUR*SE*USD",
                "ITD*01*3*2**10**30"]
        total = Decimal("0")
        for line_no, qty, uom, price, sku, desc in lines:
            segs.append(f"IT1*{line_no}*{qty}*{uom}*{price}**VP*{sku}")
            segs.append(f"PID*F****{desc}")
            total += money(Decimal(price) * qty)
        segs.append(f"TDS*{int(total * 100)}")
        segs.append(f"CTT*{len(lines)}")
        return segs

    segs = [isa("KESTRELFAST", "HARBORLINE", "260922", "0930", 4127),
            "GS*IN*KESTRELFAST*HARBORLINE*20260922*0930*4127*X*004010"]
    segs += transaction("0001", "810", body("KF-INV-88412", "20260921", "PO-2026-00412", "20260902", inv1))
    segs += transaction("0002", "810", body("KF-INV-88419", "20260922", "PO-2026-00437", "20260908", inv2))
    segs += ["GE*2*4127", "IEA*1*000004127"]
    write(os.path.join(INBOUND, "KESTREL", "KESTREL_810_20260922_0930.x12"), x12(segs))

    # A second transmission with a broken envelope: SE01 undercounts segments, so the 997 rejects it.
    bad = [isa("KESTRELFAST", "HARBORLINE", "260923", "0815", 4131),
           "GS*IN*KESTRELFAST*HARBORLINE*20260923*0815*4131*X*004010"]
    tx = transaction("0001", "810", body("KF-INV-88430", "20260923", "PO-2026-00412", "20260902",
                                        [("1", 800, "EA", "0.42", "KF-HEX-M8-40", "HEX BOLT M8X40 ZINC")]))
    tx[-1] = "SE*7*0001"
    bad += tx + ["GE*1*4131", "IEA*1*000004131"]
    write(os.path.join(ROOT, "samples", "extra", "KESTREL_810_20260923_0815_broken.x12"), x12(bad))


def brightmart():
    po1 = [
        ("1", 40, "BX", "12.50", "HL-GLV-NIT-M", "BM-778812", "NITRILE GLOVES MEDIUM BX/100"),
        ("2", 40, "BX", "12.50", "HL-GLV-NIT-L", "BM-778813", "NITRILE GLOVES LARGE BX/100"),
        ("3", 25, "CS", "58.00", "HL-TAPE-48-CLR", "BM-102233", "CARTON TAPE 48MM CLEAR CS/36"),
        ("4", 60, "BD", "32.40", "HL-BOX-RSC-18", "BM-554019", "RSC BOX 18X12X12 BUNDLE/25"),   # -6.1% vs list: warning
        ("5", 12, "RL", "14.25", "HL-LBL-4X6-500", "BM-661120", "THERMAL LABEL 4X6 ROLL"),
    ]
    po2 = [
        ("1", 200, "EA", "3.95", "HL-BRG-6204", "BM-880120", "BALL BEARING 6204-2RS"),
        ("2", 30, "RL", "44.00", "HL-STRAP-PP-12", "BM-880131", "PP STRAPPING 1/2IN COIL"),
        ("3", 80, "EA", "14.80", "HL-FLG-SS304-2", "BM-880145", "SS304 FLANGE 2IN"),
        ("4", 10, "EA", "9.99", "HL-HOSE-GRD-50", "BM-880152", "GARDEN HOSE 50FT"),              # not in catalog: error
    ]

    def body(po_no, po_date, need_by, ship_to_name, ship_to, lines):
        segs = [f"BEG*00*SA*{po_no}**{po_date}", "REF*DP*092", f"DTM*002*{need_by}",
                f"N1*ST*{ship_to_name}*92*{ship_to}"]
        for line_no, qty, uom, price, sku, buyer_part, desc in lines:
            segs.append(f"PO1*{line_no}*{qty}*{uom}*{price}**VP*{sku}*BP*{buyer_part}")
            segs.append(f"PID*F****{desc}")
        segs.append(f"CTT*{len(lines)}")
        return segs

    segs = [isa("BRIGHTMART", "HARBORLINE", "260923", "1115", 981),
            "GS*PO*BRIGHTMART*HARBORLINE*20260923*1115*981*X*004010"]
    segs += transaction("0001", "850", body("BM-4471902", "20260923", "20260930", "BRIGHTMART DC 04", "DC04", po1))
    segs += transaction("0002", "850", body("BM-4471915", "20260923", "20261002", "BRIGHTMART DC 19", "DC19", po2))  # unknown DC
    segs += ["GE*2*981", "IEA*1*000000981"]
    write(os.path.join(INBOUND, "BRIGHTMART", "BRIGHTMART_850_20260923_1115.x12"), x12(segs))


def lakeshore():
    items = [
        (1, 2400, "EA", "1.08", "LP-BOX-RSC-18", "Corrugated RSC box 18x12x12"),
        (2, 1200, "EA", "1.94", "LP-BOX-RSC-24", "Corrugated RSC box 24x18x18"),
        (3, 720, "RL", "2.35", "LP-TAPE-48", "Carton sealing tape 48mm x 100m"),
        (4, 520, "RL", "17.90", "LP-WRAP-18", "Stretch wrap 18in x 1500ft"),     # 520 > 480 ordered (+5% = 504): error
        (5, 60, "RL", "38.00", "LP-VOID-FILL", "Kraft void-fill paper roll"),
        (6, 300, "RL", "9.40", "LP-LBL-4X6", "Thermal labels 4x6, roll/500"),
    ]
    lines = []
    subtotal = Decimal("0")
    for n, qty, uom, price, sku, desc in items:
        amount = money(Decimal(price) * qty)
        subtotal += amount
        lines.append(f"""        <InvoiceDetailItem invoiceLineNumber="{n}" quantity="{qty}">
          <UnitOfMeasure>{uom}</UnitOfMeasure>
          <UnitPrice><Money currency="USD">{price}</Money></UnitPrice>
          <InvoiceDetailItemReference lineNumber="{n}">
            <ItemID><SupplierPartID>{sku}</SupplierPartID></ItemID>
            <Description xml:lang="en">{desc}</Description>
          </InvoiceDetailItemReference>
          <SubtotalAmount><Money currency="USD">{amount}</Money></SubtotalAmount>
        </InvoiceDetailItem>""")
    xml = f"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE cXML SYSTEM "http://xml.cxml.org/schemas/cXML/1.2.060/InvoiceDetail.dtd">
<cXML payloadID="20260919.1402.88712@lakeshorepkg.example" timestamp="2026-09-19T14:02:11-05:00" xml:lang="en-US">
  <Header>
    <From><Credential domain="DUNS"><Identity>LAKESHORE-PKG</Identity></Credential></From>
    <To><Credential domain="NetworkID"><Identity>HARBORLINE</Identity></Credential></To>
    <Sender><Credential domain="DUNS"><Identity>LAKESHORE-PKG</Identity></Credential><UserAgent>Lakeshore ERP 9.2</UserAgent></Sender>
  </Header>
  <Request deploymentMode="production">
    <InvoiceDetailRequest>
      <InvoiceDetailRequestHeader invoiceID="LP-26-10877" purpose="standard" operation="new" invoiceDate="2026-09-19T00:00:00-05:00">
        <InvoiceDetailHeaderIndicator/>
        <InvoiceDetailLineIndicator/>
        <InvoicePartner><Contact role="remitTo"><Name xml:lang="en">Lakeshore Packaging Co.</Name></Contact></InvoicePartner>
      </InvoiceDetailRequestHeader>
      <InvoiceDetailOrder>
        <InvoiceDetailOrderInfo>
          <OrderReference orderID="PO-2026-00455"><DocumentReference payloadID="PO-2026-00455@harborline.example"/></OrderReference>
        </InvoiceDetailOrderInfo>
{chr(10).join(lines)}
      </InvoiceDetailOrder>
      <InvoiceDetailSummary>
        <SubtotalAmount><Money currency="USD">{subtotal}</Money></SubtotalAmount>
        <Tax><Money currency="USD">0.00</Money><Description xml:lang="en">Resale certificate on file</Description></Tax>
        <NetAmount><Money currency="USD">{subtotal}</Money></NetAmount>
      </InvoiceDetailSummary>
    </InvoiceDetailRequest>
  </Request>
</cXML>
"""
    write(os.path.join(INBOUND, "LAKESHORE", "lakeshore_invoice_LP-26-10877.xml"), xml)


def sahyadri():
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice Register"
    ws["A1"] = "SAHYADRI POLYMERS PVT. LTD."
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Plot 14, MIDC Chakan Phase II, Pune 410501  |  GSTIN {SAHYADRI_GSTIN}"
    ws["A3"] = "Invoice Register - Harborline Supply Co. - September 2026"
    groups = ["Sr No", "Invoice", None, "PO", None, "Item Code", "Item Description", "HSN", "Qty", "UOM",
              "Rate (₹)", "Taxable Value (₹)", "GST (₹)", None, None, "Supplier GSTIN"]
    subs = [None, "No.", "Date", "Number", "Line", None, None, None, None, None, None, None, "CGST", "SGST", "IGST", None]
    header_fill = PatternFill("solid", fgColor="DDE7F0")
    for col, (g, s) in enumerate(zip(groups, subs), start=1):
        top = ws.cell(row=5, column=col, value=g)
        bottom = ws.cell(row=6, column=col, value=s)
        for c in (top, bottom):
            c.font = Font(bold=True)
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center")
    ws.merge_cells("B5:C5")
    ws.merge_cells("D5:E5")
    ws.merge_cells("M5:O5")
    good = SAHYADRI_GSTIN
    typo = SAHYADRI_GSTIN[:-1] + ("X" if SAHYADRI_GSTIN[-1] != "X" else "Y")
    rows = [
        # inv, date, po, po_line, sku, desc, hsn, qty, uom, rate, intra-state?, gstin, taxable override
        ("SP/26-27/0913", "15-09-2026", "HLIN-PO-2627-0118", 1, "SP-HDPE-GRN-25", "HDPE Granules Natural 25kg", "3901", 120, "BAG", 2850, True, good, None),
        ("SP/26-27/0913", "15-09-2026", "HLIN-PO-2627-0118", 2, "SP-PP-GRN-25", "PP Granules 25kg", "3902", 100, "BAG", 2640, True, good, None),
        ("SP/26-27/0913", "15-09-2026", "HLIN-PO-2627-0118", 3, "SP-MB-BLK-25", "Black Masterbatch 25kg", "3206", 20, "BAG", 4150, True, good, None),
        ("SP/26-27/0921", "18-09-2026", "HLIN-PO-2627-0118", 4, "SP-LLDPE-FLM-500", "LLDPE Film Roll 500mm", "3920", 80, "ROLL", 1980, False, good, None),
        ("SP/26-27/0921", "18-09-2026", "HLIN-PO-2627-0126", 1, "SP-HDPE-GRN-25", "HDPE Granules Natural 25kg", "3901", 100, "BAG", 2850, False, good, None),
        ("SP/26-27/0921", "18-09-2026", "HLIN-PO-2627-0126", 2, "SP-PVC-CMP-25", "PVC Compound 25kg", "3904", 60, "BAG", 3420, False, good, 202500),   # typo in value
        ("SP/26-27/0927", "22-09-2026", "HLIN-PO-2627-0126", 1, "SP-HDPE-GRN-25", "HDPE Granules Natural 25kg", "3901", 80, "BAG", 2850, True, typo, None),  # bad GSTIN
        ("SP/26-27/0927", "22-09-2026", "HLIN-PO-2627-0126", 2, "SP-PVC-CMP-25", "PVC Compound 25kg", "3904", 40, "BAG", 3420, True, typo, None),
    ]
    r = 7
    grand = Decimal("0")
    for i, (inv, date, po, po_line, sku, desc, hsn, qty, uom, rate, intra, g, override) in enumerate(rows, start=1):
        taxable = money(override if override is not None else Decimal(rate) * qty)
        gst = money(Decimal(rate) * qty * Decimal("0.18"))
        cgst = sgst = igst = None
        if intra:
            cgst = sgst = inr(gst / 2)
        else:
            igst = inr(gst)
        values = [i, inv, date, po, po_line, sku, desc, hsn, qty, uom, inr(rate), inr(taxable), cgst, sgst, igst, g]
        for col, v in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=v)
        grand += taxable
        r += 1
    ws.cell(row=r + 1, column=1, value="Grand Total")
    ws.cell(row=r + 1, column=12, value=inr(grand))
    for col, width in zip("ABCDEFGHIJKLMNOP", [7, 15, 12, 20, 6, 18, 28, 7, 6, 6, 11, 16, 11, 11, 11, 18]):
        ws.column_dimensions[col].width = width
    path = os.path.join(INBOUND, "SAHYADRI", "Sahyadri_Invoice_Register_Sep2026.xlsx")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


def narmada():
    g = NARMADA_GSTIN
    rows = [
        ("NMW/2627/0412", "03/09/2026", 1, "NMW-FLG-SS304-2", 'Flange, SS304 slip-on, 2"', "7307", 150, 685.00),
        ("NMW/2627/0412", "03/09/2026", 2, "NMW-ELB-SS304-2", 'Elbow 90°, SS304, 2"', "7307", 300, 312.00),
        ("NMW/2627/0412", "03/09/2026", 3, "NMW-TEE-SS304-2", 'Tee equal, SS304, 2"', "7307", 120, 468.00),
        None,
        ("NMW/2627/0419", "07/09/2026", 1, "NMW-FLG-SS304-3", 'Flange, SS304 slip-on, 3"', "7307", 100, 1140.00),
        ("NMW/2627/0419", "07/09/2026", 2, "NMW-PIPE-SS304-2-6M", 'Pipe SS304 2" Sch40, 6 m', "7306", 40, 9850.00),
        ("NMW/2627/0419", "07/09/2026", 2, "NMW-PIPE-SS304-2-6M", 'Pipe SS304 2" Sch40, 6 m', "7306", 40, 9850.00),   # duplicate line
        ("NMW/2627/0426", "11/09/2026", 1, "NMW-FLG-SS304-2", 'Flange, SS304 slip-on, 2"', "7307", 200, 722.00),     # +5.4% vs PO: error
    ]

    def q(s):
        return '"' + s.replace('"', '""') + '"'

    out = ["Narmada Metal Works - Invoice Lines for Harborline Supply Co.",
           q(f"Supplier GSTIN: {g}") + ",,,,,,,,,,,",
           "Inv No.,Inv Dt,PO Ref,Line,Item Code,Item Description,HSN,Qty (Nos),Rate (Rs),Amt (Rs),Tax Amt (Rs),GSTIN"]
    total = Decimal("0")
    for row in rows:
        if row is None:
            out.append(",,,,,,,,,,,")
            continue
        inv, date, line, sku, desc, hsn, qty, rate = row
        amount = money(Decimal(str(rate)) * qty)
        tax = money(amount * Decimal("0.18"))
        total += amount
        out.append(",".join([inv, date, "HLIN-PO-2627-0144", str(line), sku, q(desc), hsn, str(qty),
                             q(inr(rate)), q(inr(amount)), q(inr(tax)), g]))
    out.append(",".join(["Total", "", "", "", "", "", "", "", "", q(inr(total)), "", ""]))
    write(os.path.join(INBOUND, "NARMADA", "NMW_INV_SEP26.csv"), "\r\n".join(out) + "\r\n")


def midwest():
    def h(inv, date, po):
        return f"H{inv:<14}{date}{po:<15}MBS00017USD"

    def d(line, part, desc, qty, uom, price, amount):
        return (f"D{line:04d}{part:<20}{desc:<30}{qty:09d}{uom:<3}"
                f"{int(Decimal(price) * 10000):011d}{int(Decimal(amount) * 100):013d}")

    invoices = [
        ("MBS-INV-30118", "20260917", [
            ("MBS-6204-2RS", "BRG DEEP GROOVE 6204-2RS", 800, "EA", "2.18"),
            ("MBS-6205-2RS", "BRG DEEP GROOVE 6205-2RS", 600, "EA", "2.64"),
            ("MBS-UCP205", "PILLOW BLOCK UCP205", 120, "EA", "11.75"),
            ("MBS-SEAL-TC-35", "OIL SEAL TC 35X52X7", 1000, "EA", "0.58"),
        ]),
        ("MBS-INV-30124", "20260919", [
            ("MBS-GREASE-EP2", "GREASE EP2 CARTRIDGE 14OZ", 240, "EA", "4.95"),
            ("MBS-6204-2RS", "BRG DEEP GROOVE 6204-2RS", 400, "EA", "2.24"),    # +2.75% vs PO: warning
            ("MBS-SEAL-TC-35", "OIL SEAL TC 35X52X7", 800, "EA", "0.58"),
        ]),
    ]
    lines = []
    total = Decimal("0")
    count = 0
    for inv, date, items in invoices:
        lines.append(h(inv, date, "PO-2026-00468"))
        for n, (part, desc, qty, uom, price) in enumerate(items, start=1):
            amount = money(Decimal(price) * qty)
            total += amount
            count += 1
            lines.append(d(n, part, desc, qty, uom, price, amount))
    # The trailer claims 8 detail records but only 7 were sent: a truncated transfer the control total catches.
    lines.append(f"T{count + 1:06d}{int(total * 100):015d}")
    write(os.path.join(INBOUND, "MIDWEST", "MBS_AP_EXPORT_20260920.dat"), "\r\n".join(lines) + "\r\n")


def nimbus():
    lines = [
        (1, "NC-PCB-CTRL-V3", "Controller PCB assembly v3", 200, "PCS", "18.40"),
        (2, "NC-HARN-12P", "12-pin wiring harness, 1 m", 600, "PCS", "3.25"),
        (3, "NC-SENS-TEMP-K", "K-type thermocouple sensor", 300, "PCS", "6.10"),
        (4, "NC-ENCL-IP65", "IP65 ABS enclosure 200x150", 150, "PCS", "7.80"),
    ]
    doc = {
        "supplier": {"code": "NIMBUS", "name": "Nimbus Components Pvt. Ltd.", "gstin": NIMBUS_GSTIN},
        "invoice": {"number": "NC/EXP/26-27/0088", "date": "2026-09-20", "currency": "USD",
                    "po": "HLIN-PO-2627-0160", "exportType": "LUT (zero-rated)"},
        "lines": [{"lineNo": n, "partNo": sku, "desc": desc, "qty": qty, "uom": uom,
                   "unitPrice": float(price), "amount": float(money(Decimal(price) * qty)), "igst": 0}
                  for n, sku, desc, qty, uom, price in lines],
    }
    write(os.path.join(INBOUND, "NIMBUS", "nimbus_invoice_NC-EXP-2627-0088.json"), json.dumps(doc, indent=2) + "\n")


def coastal():
    """A supplier that is not onboarded yet: used to demo AI mapping suggestions in the Mapping Studio."""
    rows = [
        ("CC-5561", "09/04/2026", "PO-2026-00473", 1, "CC-IPA-99-1G", "Isopropyl alcohol 99%, 1 gal", 48, "EA", "18.90"),
        ("CC-5561", "09/04/2026", "PO-2026-00473", 2, "CC-DEGR-CIT-5G", "Citrus degreaser, 5 gal pail", 12, "EA", "64.50"),
        ("CC-5561", "09/04/2026", "PO-2026-00473", 3, "CC-WIPE-IND-200", "Industrial wipes, tub/200", 60, "EA", "11.25"),
        ("CC-5570", "09/11/2026", "PO-2026-00473", 1, "CC-LUB-SIL-11", "Silicone spray lubricant 11 oz", 96, "EA", "6.40"),
        ("CC-5570", "09/11/2026", "PO-2026-00473", 2, "CC-ABS-PAD-100", "Oil absorbent pads, pack/100", 20, "PK", "42.00"),
    ]
    out = ["Invoice #,Invoice Date,Customer PO,Ln,Product Code,Product Name,Qty Shipped,Unit,Price Each,Extended Price,Sales Tax,Currency"]
    for inv, date, po, ln, code, name, qty, unit, price in rows:
        ext = money(Decimal(price) * qty)
        out.append(",".join([inv, date, po, str(ln), code, f'"{name}"', str(qty), unit, f"${price}", f'"${ext:,}"', "$0.00", "USD"]))
    write(os.path.join(ROOT, "samples", "extra", "coastal_chem_invoice_lines.csv"), "\n".join(out) + "\n")


# ------------------------------------------------------------------ reference data outputs


def reference_json():
    data = {
        "po": [{"poNumber": po, "supplierId": sup, "currency": cur, "poDate": date, "status": "OPEN"}
               for po, sup, cur, date, _ in PURCHASE_ORDERS],
        "po-line": [{"poNumber": po, "sku": sku, "description": desc, "orderedQty": qty, "invoicedQty": 0,
                     "unitPrice": price, "uom": uom}
                    for po, _, _, _, lines in PURCHASE_ORDERS for sku, desc, qty, price, uom in lines],
        "catalog": [{"sku": sku, "description": desc, "uom": uom, "listPrice": price, "currency": "USD"}
                    for sku, desc, uom, price in CATALOG],
        "ship-to": [{"customerId": c, "shipToCode": code, "name": name, "city": city, "state": st}
                    for c, code, name, city, st in SHIP_TO],
        "fx": [{"currency": c, "rate": r} for c, r in FX],
    }
    write(os.path.join(BERTH9, "reference", "reference-data.json"), json.dumps(data, indent=2) + "\n")
    write(os.path.join(BERTH9, "partners.json"), json.dumps(PARTNERS, indent=2, ensure_ascii=False) + "\n")


def sql_seed():
    def s(v):
        return "'" + str(v).replace("'", "''") + "'"

    out = ["-- Demo reference data for Harborline Supply Co. (fictional). Generated by samples/generate.py.", ""]
    for p in PARTNERS:
        out.append("INSERT INTO partner (id, name, country, city, role, channel, format, document, schema_name, currency, "
                   "gstin, edi_id, locale, expected_by, api_key, api_secret) VALUES ("
                   + ", ".join([s(p["id"]), s(p["name"]), s(p["country"]), s(p["city"]), s(p["role"]), s(p["channel"]),
                                s(p["format"]), s(p["document"]), s(p["schema"]), s(p["currency"]),
                                s(p["gstin"]) if p.get("gstin") else "NULL", s(p["ediId"]) if p.get("ediId") else "NULL",
                                s(p["locale"]), s(p["expectedBy"]),
                                s("pk_" + p["id"].lower()), s("demo-secret-" + p["id"].lower())]) + ");")
    out.append("")
    for po, sup, cur, date, lines in PURCHASE_ORDERS:
        out.append(f"INSERT INTO purchase_order (po_number, supplier_id, currency, po_date, status) VALUES "
                   f"({s(po)}, {s(sup)}, {s(cur)}, DATE {s(date)}, 'OPEN');")
        for sku, desc, qty, price, uom in lines:
            out.append(f"INSERT INTO po_line (po_number, sku, description, ordered_qty, invoiced_qty, unit_price, uom) VALUES "
                       f"({s(po)}, {s(sku)}, {s(desc)}, {qty}, 0, {price}, {s(uom)});")
    out.append("")
    for sku, desc, uom, price in CATALOG:
        out.append(f"INSERT INTO product (sku, description, uom, list_price, currency) VALUES ({s(sku)}, {s(desc)}, {s(uom)}, {price}, 'USD');")
    out.append("")
    for c, code, name, city, st in SHIP_TO:
        out.append(f"INSERT INTO ship_to (customer_id, ship_to_code, name, city, state) VALUES ({s(c)}, {s(code)}, {s(name)}, {s(city)}, {s(st)});")
    out.append("")
    for c, r in FX:
        out.append(f"INSERT INTO fx_rate (currency, usd_rate) VALUES ({s(c)}, {r});")
    write(os.path.join(RESOURCES, "db", "migration", "V2__demo_reference_data.sql"), "\n".join(out) + "\n")


if __name__ == "__main__":
    kestrel()
    brightmart()
    lakeshore()
    sahyadri()
    narmada()
    midwest()
    nimbus()
    coastal()
    reference_json()
    sql_seed()
    import shutil
    demo = os.path.join(RESOURCES, "demo")
    shutil.rmtree(demo, ignore_errors=True)
    shutil.copytree(INBOUND, demo)
    print("GSTINs:", SAHYADRI_GSTIN, NARMADA_GSTIN, NIMBUS_GSTIN, BUYER_GSTIN)
    print("samples written to", INBOUND)
